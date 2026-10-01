from __future__ import annotations

import re
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt

from fpdf import FPDF

from backend.config import PROJECT_ROOT
from backend.utils.text_utils import sanitize_text


LOGO_PATH = (
    PROJECT_ROOT
    / "assets"
    / "logo.png"
)


def get_lines(text: str) -> list[str]:

    clean_text = sanitize_text(text)

    return [
        line.strip()
        for line in clean_text.splitlines()
        if line.strip()
    ]


def is_heading(line: str) -> bool:

    upper = line.upper()

    if upper.startswith("ARTICLE "):
        return True

    if upper.startswith("SECTION "):
        return True

    if re.match(
        r"^[1-9][0-9]?\.\s",
        line
    ):
        return True

    headings = {
        "PARTIES",
        "RECITALS",
        "DEFINITIONS",
        "CONFIDENTIALITY",
        "PAYMENT",
        "TERM AND TERMINATION",
        "GOVERNING LAW",
        "SIGNATURES",
        "SIGNATURE",
    }

    return upper in headings


# ============================================================
# TXT
# ============================================================

def format_txt(text: str) -> bytes:

    clean_text = sanitize_text(text)

    return clean_text.encode(
        "utf-8"
    )


# ============================================================
# DOCX
# ============================================================

def format_docx(
    text: str,
    document_type: str
) -> bytes:

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal_style = document.styles["Normal"]

    normal_style.font.name = (
        "Times New Roman"
    )

    normal_style.font.size = Pt(11)

    # Footer
    footer = section.footer.paragraphs[0]

    footer.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    footer_run = footer.add_run(
        "Generated with LegalEase • "
        "Review before use"
    )

    footer_run.font.size = Pt(8)

    # Optional logo
    if LOGO_PATH.exists():

        logo_paragraph = document.add_paragraph()

        logo_paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        logo_paragraph.add_run().add_picture(
            str(LOGO_PATH),
            width=Inches(1.2)
        )

    # Title
    title = document.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    title_run = title.add_run(
        document_type.upper()
    )

    title_run.bold = True
    title_run.font.name = (
        "Times New Roman"
    )
    title_run.font.size = Pt(16)

    lines = get_lines(text)

    bullet_terms = []

    for line in lines:

        if line.upper() == document_type.upper():
            continue

        # Bullet
        if line.startswith(
            ("-", "•", "*")
        ):

            clean_bullet = re.sub(
                r"^[-•*]\s*",
                "",
                line
            )

            bullet_terms.append(
                clean_bullet
            )

            paragraph = document.add_paragraph(
                style="List Bullet"
            )

            paragraph.add_run(
                clean_bullet
            )

        # Heading
        elif is_heading(line):

            paragraph = document.add_paragraph()

            run = paragraph.add_run(
                line
            )

            run.bold = True
            run.font.size = Pt(12)

            paragraph.paragraph_format.space_before = (
                Pt(8)
            )

            paragraph.paragraph_format.space_after = (
                Pt(4)
            )

        # Normal paragraph
        else:

            paragraph = document.add_paragraph(
                line
            )

            paragraph.paragraph_format.line_spacing = (
                1.15
            )

            paragraph.paragraph_format.space_after = (
                Pt(6)
            )

    # Terms table
    if bullet_terms:

        document.add_page_break()

        heading = document.add_paragraph()

        heading_run = heading.add_run(
            "Key Terms Summary"
        )

        heading_run.bold = True

        table = document.add_table(
            rows=1,
            cols=2
        )

        table.style = "Table Grid"

        table.alignment = (
            WD_TABLE_ALIGNMENT.CENTER
        )

        header_cells = table.rows[0].cells

        header_cells[0].text = "No."
        header_cells[1].text = "Term"

        for index, term in enumerate(
            bullet_terms,
            start=1
        ):

            cells = table.add_row().cells

            cells[0].text = str(index)
            cells[1].text = term

    output = BytesIO()

    document.save(output)

    return output.getvalue()


# ============================================================
# PDF
# ============================================================

class LegalEasePDF(FPDF):

    def __init__(
        self,
        document_type: str
    ):

        super().__init__()

        self.document_type = (
            document_type
        )

        self.set_auto_page_break(
            auto=True,
            margin=18
        )

    def header(self):

        if LOGO_PATH.exists():

            try:

                self.image(
                    str(LOGO_PATH),
                    x=92,
                    y=8,
                    w=25
                )

                self.ln(16)

            except Exception:

                self.ln(5)

        else:

            self.ln(5)

        self.set_font(
            "Helvetica",
            "B",
            11
        )

        self.cell(
            0,
            7,
            self.document_type.upper(),
            align="C"
        )

        self.ln(8)

    def footer(self):

        self.set_y(-14)

        self.set_font(
            "Helvetica",
            "",
            8
        )

        self.cell(
            0,
            8,
            "Generated with LegalEase - "
            "Review before use",
            align="C"
        )


def format_pdf(
    text: str,
    document_type: str
) -> bytes:

    pdf = LegalEasePDF(
        document_type
    )

    pdf.add_page()

    lines = get_lines(text)

    for line in lines:

        if line.upper() == document_type.upper():
            continue

        if is_heading(line):

            pdf.set_font(
                "Helvetica",
                "B",
                11
            )

            pdf.ln(2)

            pdf.multi_cell(
                0,
                6,
                line
            )

        else:

            pdf.set_font(
                "Helvetica",
                "",
                10
            )

            if line.startswith(
                ("-", "•", "*")
            ):

                line = (
                    "- "
                    + re.sub(
                        r"^[-•*]\s*",
                        "",
                        line
                    )
                )

            pdf.multi_cell(
                0,
                5.5,
                line
            )

            pdf.ln(1)

    return bytes(
        pdf.output()
    )